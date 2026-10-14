"""Plan 007: the openpyxl adapters. The template writer (headers, hidden lists, dropdowns), the
sheet reader (headers by key, empty rows, row limit, unreadable files, formulas) and the round
trip between them. They need no service; they sit with the integration tests because the plan
files them there."""

import zipfile
from datetime import date, datetime
from io import BytesIO
from typing import Any

import pytest
from openpyxl import Workbook, load_workbook  # type: ignore[import-untyped]
from openpyxl.styles import Font  # type: ignore[import-untyped]

from fragancia_api.modules.catalog.domain.import_cells import COLUMNS, header_key
from fragancia_api.modules.catalog.infrastructure import xlsx_sheets
from fragancia_api.modules.catalog.infrastructure.xlsx_sheets import (
    XLSX_MAX_ENTRIES,
    XLSX_MAX_UNCOMPRESSED_BYTES,
    OpenpyxlSheetReader,
    OpenpyxlTemplateWriter,
)
from fragancia_api.shared.kernel import Err
from tests.unit.catalog.perfume_support import unwrap_ok

pytestmark = pytest.mark.integration

FAMILIES = ["Aromática", "Cítrica", "Floral"]
CONCENTRATIONS = ["Eau de Parfum", "Eau de Toilette"]
MAX_ROWS = 2000


def _workbook(rows: list[list[Any]], *, header: list[Any] | None = None) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(header if header is not None else list(COLUMNS))
    for values in rows:
        sheet.append(values)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _patched(data: bytes, replacements: dict[str, str]) -> bytes:
    """The same workbook with text replaced inside its first sheet's XML."""
    output = BytesIO()
    with zipfile.ZipFile(BytesIO(data)) as source, zipfile.ZipFile(output, "w") as target:
        for item in source.infolist():
            content = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                text = content.decode()
                for old, new in replacements.items():
                    assert old in text, old
                    text = text.replace(old, new)
                content = text.encode()
            target.writestr(item, content)
    return output.getvalue()


async def _read(data: bytes, max_rows: int = MAX_ROWS) -> Any:
    return await OpenpyxlSheetReader().read(data, max_rows=max_rows)


def _code(result: Any) -> str:
    assert isinstance(result, Err), result
    return str(result.error.code)


# --- the template -----------------------------------------------------------------------------


async def _template(
    families: list[str] = FAMILIES, concentrations: list[str] = CONCENTRATIONS
) -> Any:
    data = await OpenpyxlTemplateWriter().build(families, concentrations)
    return load_workbook(BytesIO(data))


async def test_the_template_has_the_catalogo_sheet_with_the_21_headers_in_order() -> None:
    workbook = await _template()

    sheet = workbook["Catálogo"]
    assert workbook.worksheets[0].title == "Catálogo"  # the reader reads the first sheet
    assert [cell.value for cell in sheet[1]] == list(COLUMNS)
    assert all(cell.font.bold for cell in sheet[1])
    assert sheet.freeze_panes == "A2"
    assert sheet.max_row == 1  # no data rows
    assert all(sheet.column_dimensions[letter].width >= 14 for letter in "ABCDEFGHIJKLMNOPQRSTU")


async def test_the_template_has_existencias_after_the_max_lead_time_with_no_dropdown() -> None:
    workbook = await _template()

    sheet = workbook["Catálogo"]
    headers = [cell.value for cell in sheet[1]]
    index = headers.index("Entrega máx. (días)")
    assert headers[index + 1] == "Existencias"
    assert sheet.cell(row=1, column=index + 2).font.bold
    letter = chr(ord("A") + index + 1)
    assert sheet.column_dimensions[letter].width >= 14
    assert all(
        letter not in str(validation.sqref)[0]
        for validation in sheet.data_validations.dataValidation
    )


async def test_a_count_written_under_existencias_is_read_back_by_the_reader() -> None:
    data = await OpenpyxlTemplateWriter().build(FAMILIES, CONCENTRATIONS)
    workbook = load_workbook(BytesIO(data))
    values_by_header = {"Marca": "Versace", "Nombre": "Eros", "Existencias": 12}
    workbook["Catálogo"].append([values_by_header.get(column) for column in COLUMNS])
    buffer = BytesIO()
    workbook.save(buffer)

    sheet = unwrap_ok(await _read(buffer.getvalue()))

    [(_, values)] = sheet.rows
    assert values[header_key("Existencias")] == 12


async def test_the_template_has_a_hidden_lists_sheet_with_every_dropdown_value() -> None:
    workbook = await _template()

    lists = workbook["Listas"]
    assert lists.sheet_state == "hidden"
    assert workbook["Catálogo"].sheet_state == "visible"

    def column(letter: str) -> list[Any]:
        return [cell.value for cell in lists[letter] if cell.value is not None]

    assert column("A") == FAMILIES
    assert column("B") == CONCENTRATIONS
    assert column("C") == ["Mujer", "Hombre", "Unisex"]
    assert column("D") == ["En existencia", "Sobre pedido"]
    assert column("E") == ["Sí", "No"]


async def test_the_template_puts_a_list_dropdown_on_the_five_columns_for_rows_2_to_2001() -> None:
    workbook = await _template()

    validations = {
        str(validation.sqref): validation
        for validation in workbook["Catálogo"].data_validations.dataValidation
    }
    letter = {name: chr(ord("A") + COLUMNS.index(name)) for name in COLUMNS}
    expected = {
        "Familia olfativa": f"=Listas!$A$1:$A${len(FAMILIES)}",
        "Concentración": f"=Listas!$B$1:$B${len(CONCENTRATIONS)}",
        "Género": "=Listas!$C$1:$C$3",
        "Disponibilidad": "=Listas!$D$1:$D$2",
        "Publicado": "=Listas!$E$1:$E$2",
    }
    assert set(validations) == {f"{letter[name]}2:{letter[name]}2001" for name in expected}
    for name, formula in expected.items():
        validation = validations[f"{letter[name]}2:{letter[name]}2001"]
        assert (validation.type, validation.formula1) == ("list", formula)
        assert validation.allow_blank is True


async def test_the_template_skips_the_dropdown_of_an_empty_list() -> None:
    workbook = await _template(families=[], concentrations=[])

    sqrefs = {str(v.sqref) for v in workbook["Catálogo"].data_validations.dataValidation}
    family = chr(ord("A") + COLUMNS.index("Familia olfativa"))
    concentration = chr(ord("A") + COLUMNS.index("Concentración"))
    assert f"{family}2:{family}2001" not in sqrefs
    assert f"{concentration}2:{concentration}2001" not in sqrefs
    assert len(sqrefs) == 3  # gender, availability and published remain


async def test_an_empty_template_is_read_as_an_empty_file() -> None:
    data = await OpenpyxlTemplateWriter().build(FAMILIES, CONCENTRATIONS)

    result = await _read(data)

    assert _code(result) == "CATALOG_IMPORT_EMPTY"


async def test_a_filled_template_round_trips_through_the_reader() -> None:
    data = await OpenpyxlTemplateWriter().build(FAMILIES, CONCENTRATIONS)
    workbook = load_workbook(BytesIO(data))
    workbook["Catálogo"].append(
        ["Versace", "Eros", "Eau de Toilette", "Hombre", "Aromática", 100, 2500.5]
    )
    buffer = BytesIO()
    workbook.save(buffer)

    sheet = unwrap_ok(await _read(buffer.getvalue()))

    assert sheet.headers == [header_key(c) for c in COLUMNS]
    assert sheet.duplicated == []
    [(number, cells)] = sheet.rows
    assert number == 2
    assert cells["marca"] == "Versace"
    assert cells["concentracion"] == "Eau de Toilette"
    assert (cells["ml"], cells["precio"]) == (100, 2500.5)


# --- the reader: headers ----------------------------------------------------------------------


async def test_headers_match_by_key_ignoring_case_accents_and_punctuation() -> None:
    data = _workbook([["Versace", 100]], header=["MARCA", "  ml "])

    sheet = unwrap_ok(await _read(data))

    assert sheet.headers == ["marca", "ml"]
    assert sheet.rows == [(2, {"marca": "Versace", "ml": 100})]


async def test_unknown_headers_are_kept_by_key_so_the_validator_ignores_them() -> None:
    data = _workbook([["a", "b"]], header=["Marca", "Comentarios"])

    sheet = unwrap_ok(await _read(data))

    assert sheet.rows == [(2, {"marca": "a", "comentarios": "b"})]


async def test_blank_header_cells_are_skipped_and_their_columns_dropped() -> None:
    data = _workbook([["a", "dropped", "c"]], header=["Marca", None, "ML"])

    sheet = unwrap_ok(await _read(data))

    assert sheet.headers == ["marca", "ml"]
    assert sheet.rows == [(2, {"marca": "a", "ml": "c"})]


async def test_a_header_appearing_twice_is_reported_once_and_the_first_column_wins() -> None:
    data = _workbook([["first", "second", "third"]], header=["Precio", "precio", "PRECIO"])

    sheet = unwrap_ok(await _read(data))

    assert sheet.headers == ["precio"]
    assert sheet.duplicated == ["precio"]
    assert sheet.rows == [(2, {"precio": "first"})]


async def test_only_the_first_thirty_columns_are_read() -> None:
    header = [f"Col {n}" for n in range(1, 32)]
    data = _workbook([list(range(1, 32))], header=header)

    sheet = unwrap_ok(await _read(data))

    assert len(sheet.headers) == 30
    assert "col-31" not in sheet.headers
    assert sheet.rows[0][1]["col-30"] == 30


async def test_only_the_first_worksheet_is_read() -> None:
    workbook = Workbook()
    first = workbook.active
    first.append(["Marca"])
    first.append(["from the first"])
    second = workbook.create_sheet("Otra")
    second.append(["Marca"])
    second.append(["from the second"])
    buffer = BytesIO()
    workbook.save(buffer)

    sheet = unwrap_ok(await _read(buffer.getvalue()))

    assert sheet.rows == [(2, {"marca": "from the first"})]


# --- the reader: rows -------------------------------------------------------------------------


async def test_fully_empty_rows_are_skipped_but_row_numbers_stay_the_excel_ones() -> None:
    data = _workbook([["a"], [None], ["  "], ["d"], [None]], header=["Marca"])

    sheet = unwrap_ok(await _read(data))

    assert [number for number, _ in sheet.rows] == [2, 5]
    assert [cells["marca"] for _, cells in sheet.rows] == ["a", "d"]


async def test_cell_types_are_kept_numbers_dates_and_text() -> None:
    data = _workbook(
        [[100, 2500.5, "texto", datetime(2026, 11, 1), True, None]],
        header=["ML", "Precio", "Nombre", "Oferta desde", "Publicado", "Descripción"],
    )

    sheet = unwrap_ok(await _read(data))

    [(_, cells)] = sheet.rows
    assert cells["ml"] == 100
    assert cells["precio"] == 2500.5
    assert cells["nombre"] == "texto"
    assert cells["oferta-desde"] == datetime(2026, 11, 1)
    assert cells["publicado"] is True
    assert cells["descripcion"] is None


async def test_a_date_cell_is_read_as_a_datetime_at_midnight() -> None:
    data = _workbook([[date(2026, 11, 30)]], header=["Oferta hasta"])

    sheet = unwrap_ok(await _read(data))

    value = sheet.rows[0][1]["oferta-hasta"]
    assert isinstance(value, datetime)
    assert value == datetime(2026, 11, 30, 0, 0)


async def test_a_short_row_has_none_for_the_missing_trailing_cells() -> None:
    data = _workbook([["only brand"]], header=["Marca", "Nombre"])

    sheet = unwrap_ok(await _read(data))

    assert sheet.rows == [(2, {"marca": "only brand", "nombre": None})]


async def test_a_wrong_stored_dimension_does_not_hide_rows() -> None:
    data = _patched(
        _workbook([["a"], ["b"], ["c"]], header=["Marca"]),
        {'<dimension ref="A1:A4" />': '<dimension ref="A1:A1" />'},
    )

    sheet = unwrap_ok(await _read(data))

    assert [cells["marca"] for _, cells in sheet.rows] == ["a", "b", "c"]


async def test_a_header_only_file_is_empty() -> None:
    assert _code(await _read(_workbook([]))) == "CATALOG_IMPORT_EMPTY"


async def test_rows_with_only_blanks_are_empty() -> None:
    assert (
        _code(await _read(_workbook([[None], ["   "]], header=["Marca"]))) == "CATALOG_IMPORT_EMPTY"
    )


# --- the reader: the row limit ----------------------------------------------------------------


async def test_exactly_max_rows_data_rows_are_accepted() -> None:
    data = _workbook([[n] for n in range(3)], header=["ML"])

    sheet = unwrap_ok(await _read(data, max_rows=3))

    assert len(sheet.rows) == 3


async def test_one_more_than_max_rows_is_too_many_rows_with_the_limit_in_the_error() -> None:
    data = _workbook([[n] for n in range(4)], header=["ML"])

    result = await _read(data, max_rows=3)

    assert _code(result) == "CATALOG_IMPORT_TOO_MANY_ROWS"
    assert isinstance(result, Err)
    assert result.error.message == "The file has too many data rows"
    assert result.error.row is None
    assert result.error.column is None


async def test_blank_rows_do_not_count_toward_the_limit() -> None:
    data = _workbook([[1], [None], [2], [None], [None], [3]], header=["ML"])

    sheet = unwrap_ok(await _read(data, max_rows=3))

    assert [number for number, _ in sheet.rows] == [2, 4, 7]


async def test_the_default_limit_of_2000_rows_accepts_2000_and_refuses_2001() -> None:
    ok = unwrap_ok(await _read(_workbook([[n + 1] for n in range(2000)], header=["ML"])))
    refused = await _read(_workbook([[n + 1] for n in range(2001)], header=["ML"]))

    assert len(ok.rows) == 2000
    assert _code(refused) == "CATALOG_IMPORT_TOO_MANY_ROWS"


# --- the reader: files that are not workbooks -------------------------------------------------


@pytest.mark.parametrize(
    "data",
    [
        b"",
        b"definitely not a workbook",
        b"%PDF-1.7 pretend",
        b"PK\x03\x04 but not really a zip",
        b"\x00" * 64,
    ],
)
async def test_bytes_that_are_not_an_xlsx_are_an_unreadable_file(data: bytes) -> None:
    result = await _read(data)

    assert _code(result) == "CATALOG_IMPORT_FILE_UNREADABLE"
    assert isinstance(result, Err)
    assert (result.error.row, result.error.column) == (None, None)


async def test_a_zip_that_is_not_a_workbook_is_unreadable() -> None:
    buffer = BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("hello.txt", "just a text file in a zip")

    assert _code(await _read(buffer.getvalue())) == "CATALOG_IMPORT_FILE_UNREADABLE"


async def test_a_truncated_workbook_is_unreadable() -> None:
    data = _workbook([["a"]], header=["Marca"])

    assert _code(await _read(data[: len(data) // 2])) == "CATALOG_IMPORT_FILE_UNREADABLE"


async def test_a_csv_renamed_to_xlsx_is_unreadable() -> None:
    assert _code(await _read(b"Marca,ML\nVersace,100\n")) == "CATALOG_IMPORT_FILE_UNREADABLE"


# --- the reader: formulas ---------------------------------------------------------------------


async def test_a_formula_is_read_as_its_cached_value_never_as_formula_text() -> None:
    data = _patched(
        _workbook([["=1+1"]], header=["ML"]),
        {"<f>1+1</f><v />": "<f>1+1</f><v>2</v>"},
    )

    sheet = unwrap_ok(await _read(data))

    assert sheet.rows == [(2, {"ml": 2})]


async def test_a_formula_without_a_cached_value_is_empty_not_the_formula_text() -> None:
    data = _workbook([['=HYPERLINK("http://evil.test","x")'], ["real"]], header=["Marca"])

    sheet = unwrap_ok(await _read(data))

    assert sheet.rows == [(3, {"marca": "real"})] or sheet.rows == [
        (2, {"marca": None}),
        (3, {"marca": "real"}),
    ]
    assert "evil.test" not in str(sheet.rows)


# --- the reader: bounded expansion (repair round 1, M2) ----------------------------------------


def _with_entries(data: bytes, entries: dict[str, bytes]) -> bytes:
    """The same workbook with extra ZIP entries (deflated, so big ones stay small on disk)."""
    output = BytesIO()
    with (
        zipfile.ZipFile(BytesIO(data)) as source,
        zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target,
    ):
        for item in source.infolist():
            target.writestr(item, source.read(item.filename))
        for name, content in entries.items():
            target.writestr(name, content)
    return output.getvalue()


def _entry_count(data: bytes) -> int:
    with zipfile.ZipFile(BytesIO(data)) as archive:
        return len(archive.infolist())


GOOD = _workbook([["Versace"]], header=["Marca"])


async def test_a_workbook_with_exactly_the_maximum_entries_is_read() -> None:
    padding = {f"padding/{n}.txt": b"x" for n in range(XLSX_MAX_ENTRIES - _entry_count(GOOD))}
    data = _with_entries(GOOD, padding)
    assert _entry_count(data) == XLSX_MAX_ENTRIES

    assert unwrap_ok(await _read(data)).rows == [(2, {"marca": "Versace"})]


async def test_a_workbook_with_one_entry_over_the_maximum_is_unreadable() -> None:
    padding = {f"padding/{n}.txt": b"x" for n in range(XLSX_MAX_ENTRIES - _entry_count(GOOD) + 1)}
    data = _with_entries(GOOD, padding)
    assert _entry_count(data) == XLSX_MAX_ENTRIES + 1

    assert _code(await _read(data)) == "CATALOG_IMPORT_FILE_UNREADABLE"


async def test_an_entry_declaring_more_than_the_limit_is_unreadable_before_it_is_parsed() -> None:
    bomb = b"\x00" * (XLSX_MAX_UNCOMPRESSED_BYTES + 1)
    data = _with_entries(GOOD, {"xl/bomb.bin": bomb})
    assert len(data) < 1_000_000  # it deflates to almost nothing, like a real bomb

    assert _code(await _read(data)) == "CATALOG_IMPORT_FILE_UNREADABLE"


async def test_entries_each_within_the_limit_but_together_over_it_are_unreadable() -> None:
    half = b"\x00" * (XLSX_MAX_UNCOMPRESSED_BYTES // 2 + 1)
    data = _with_entries(GOOD, {"xl/a.bin": half, "xl/b.bin": half})

    assert _code(await _read(data)) == "CATALOG_IMPORT_FILE_UNREADABLE"


async def test_a_big_but_in_limit_workbook_is_still_read() -> None:
    data = _with_entries(GOOD, {"xl/ok.bin": b"\x00" * 20_000_000})

    assert unwrap_ok(await _read(data)).rows == [(2, {"marca": "Versace"})]


# --- the reader: blank rows count toward the scan limit (repair round 1, M2) --------------------


async def test_blank_rows_up_to_the_scan_limit_are_scanned_and_skipped(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    # scan limit 10: exactly 10 physical rows after the header are allowed, blank or not.
    monkeypatch.setattr(xlsx_sheets, "XLSX_MAX_SCANNED_ROWS", 10)
    data = _workbook([[1], [2], [3]] + [[None]] * 5 + [[4]] + [[None]], header=["ML"])

    sheet = unwrap_ok(await _read(data, max_rows=5))

    assert [number for number, _ in sheet.rows] == [2, 3, 4, 10]


async def test_one_blank_row_past_the_scan_limit_is_too_many_rows(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(xlsx_sheets, "XLSX_MAX_SCANNED_ROWS", 10)
    data = _workbook([[1], [2], [3]] + [[None]] * 7 + [[4]], header=["ML"])

    result = await _read(data, max_rows=5)  # 11 physical rows after the header

    assert _code(result) == "CATALOG_IMPORT_TOO_MANY_ROWS"


async def test_a_sheet_of_whitespace_rows_stops_at_the_scan_limit(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(xlsx_sheets, "XLSX_MAX_SCANNED_ROWS", 10)
    data = _workbook([["   "]] * 40, header=["ML"])

    assert _code(await _read(data, max_rows=5)) == "CATALOG_IMPORT_TOO_MANY_ROWS"


async def test_formatted_empty_cells_down_to_row_5000_do_not_hide_the_three_data_rows() -> None:
    # Repair round 2: styled but empty cells make the sheet declare thousands of rows.
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Marca"])
    for brand in ("a", "b", "c"):
        sheet.append([brand])
    for row in range(5, 5001):
        sheet.cell(row=row, column=1).font = Font(bold=True)
    buffer = BytesIO()
    workbook.save(buffer)
    data = buffer.getvalue()
    with zipfile.ZipFile(BytesIO(data)) as archive:
        assert 'r="A5000"' in archive.read("xl/worksheets/sheet1.xml").decode()

    result = unwrap_ok(await _read(data))

    assert result.rows == [(2, {"marca": "a"}), (3, {"marca": "b"}), (4, {"marca": "c"})]
