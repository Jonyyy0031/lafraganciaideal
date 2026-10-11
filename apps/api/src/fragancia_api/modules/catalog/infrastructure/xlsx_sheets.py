"""The import sheet reader and the template writer (openpyxl, which blocks: both run in a
worker thread through `asyncio.to_thread`)."""

import asyncio
from io import BytesIO
from typing import Any

import structlog
from openpyxl import Workbook, load_workbook  # type: ignore[import-untyped]  # no stubs
from openpyxl.styles import Font  # type: ignore[import-untyped]
from openpyxl.utils import get_column_letter  # type: ignore[import-untyped]
from openpyxl.worksheet.datavalidation import DataValidation  # type: ignore[import-untyped]

from fragancia_api.modules.catalog.application.ports import SheetRows
from fragancia_api.modules.catalog.domain.catalog_import import ImportIssue
from fragancia_api.modules.catalog.domain.errors import (
    ImportEmpty,
    ImportFileUnreadable,
    ImportTooManyRows,
)
from fragancia_api.modules.catalog.domain.import_cells import COLUMNS, Cell, header_key
from fragancia_api.shared.kernel import Err, Ok, Result

log = structlog.get_logger(__name__)

MAX_COLUMNS = 30
TEMPLATE_ROWS = 2000  # rows 2-2001 get the dropdowns
_DROPDOWNS: dict[str, str] = {
    # column → its list in the hidden sheet
    "Familia olfativa": "A",
    "Concentración": "B",
    "Género": "C",
    "Disponibilidad": "D",
    "Publicado": "E",
}
_GENDERS = ["Mujer", "Hombre", "Unisex"]
_AVAILABILITIES = ["En existencia", "Sobre pedido"]
_PUBLISHED = ["Sí", "No"]


def _is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and value.strip() == "")


def _file_issue(error: ImportFileUnreadable | ImportEmpty | ImportTooManyRows) -> ImportIssue:
    return ImportIssue.of(error, row=None, column=None)


class OpenpyxlSheetReader:
    """Reads the first worksheet: the header from row 1, then the non-empty data rows."""

    async def read(self, data: bytes, *, max_rows: int) -> Result[SheetRows, ImportIssue]:
        return await asyncio.to_thread(self._read, data, max_rows)

    def _read(self, data: bytes, max_rows: int) -> Result[SheetRows, ImportIssue]:
        workbook = None
        try:
            workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
            sheet = workbook.worksheets[0]
            sheet.reset_dimensions()  # some writers store a wrong size: read every row
            return self._rows(sheet, max_rows)
        except Exception as error:  # zipfile and openpyxl raise many kinds on a bad file
            log.info("catalog_import_file_unreadable", error=type(error).__name__)
            return Err(_file_issue(ImportFileUnreadable()))
        finally:
            if workbook is not None:
                workbook.close()

    @staticmethod
    def _rows(sheet: Any, max_rows: int) -> Result[SheetRows, ImportIssue]:
        header = next(
            sheet.iter_rows(min_row=1, max_row=1, max_col=MAX_COLUMNS, values_only=True), ()
        )
        keys_by_index: dict[int, str] = {}
        headers: list[str] = []
        duplicated: list[str] = []
        for index, value in enumerate(header):
            if _is_blank(value):
                continue
            key = header_key(str(value))
            if not key:
                continue
            if key in headers:
                if key not in duplicated:
                    duplicated.append(key)
                continue
            headers.append(key)
            keys_by_index[index] = key

        rows: list[tuple[int, dict[str, Cell]]] = []
        for number, values in enumerate(
            sheet.iter_rows(min_row=2, max_col=MAX_COLUMNS, values_only=True), start=2
        ):
            if all(_is_blank(value) for value in values):
                continue
            if len(rows) == max_rows:  # this is row max_rows + 1
                return Err(_file_issue(ImportTooManyRows(details={"max_rows": max_rows})))
            rows.append(
                (
                    number,
                    {
                        key: values[index]
                        for index, key in keys_by_index.items()
                        if index < len(values)
                    },
                )
            )
        if not rows:
            return Err(_file_issue(ImportEmpty()))
        return Ok(SheetRows(rows=rows, headers=headers, duplicated=duplicated))


class OpenpyxlTemplateWriter:
    """The import template: the `Catálogo` sheet with the headers and dropdowns, and a hidden
    `Listas` sheet with the dropdown values."""

    async def build(self, families: list[str], concentrations: list[str]) -> bytes:
        return await asyncio.to_thread(self._build, families, concentrations)

    @staticmethod
    def _build(families: list[str], concentrations: list[str]) -> bytes:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Catálogo"
        bold = Font(bold=True)
        for index, column in enumerate(COLUMNS, start=1):
            cell = sheet.cell(row=1, column=index, value=column)
            cell.font = bold
            sheet.column_dimensions[get_column_letter(index)].width = max(14, len(column) + 4)
        sheet.freeze_panes = "A2"

        lists = workbook.create_sheet("Listas")
        lists.sheet_state = "hidden"
        values_by_letter = {
            "A": families,
            "B": concentrations,
            "C": _GENDERS,
            "D": _AVAILABILITIES,
            "E": _PUBLISHED,
        }
        for letter, values in values_by_letter.items():
            for row, value in enumerate(values, start=1):
                lists[f"{letter}{row}"] = value

        for column, letter in _DROPDOWNS.items():
            count = len(values_by_letter[letter])
            if count == 0:
                continue  # nothing to offer (no active families or concentrations)
            validation = DataValidation(
                type="list", formula1=f"=Listas!${letter}$1:${letter}${count}", allow_blank=True
            )
            target = get_column_letter(COLUMNS.index(column) + 1)
            validation.add(f"{target}2:{target}{TEMPLATE_ROWS + 1}")
            sheet.add_data_validation(validation)

        output = BytesIO()
        workbook.save(output)
        return output.getvalue()
